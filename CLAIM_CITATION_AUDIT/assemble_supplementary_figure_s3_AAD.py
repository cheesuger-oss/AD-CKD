#!/usr/bin/env python3
"""Build Supplementary Figure S3 from the frozen Supplementary Table S2.

This deliberately plots frozen ORA and GSEA summary statistics only. The
gene-level ranked lists needed to reproduce running-enrichment curves are not
present in the submission archive, so no curve is reconstructed or imputed.
"""

from pathlib import Path
import math
import textwrap
import json
import hashlib

import matplotlib.pyplot as plt
from matplotlib.colors import Normalize
from matplotlib.cm import ScalarMappable
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent / "aad_revision_candidate" / "figures"
BOOK = Path(__file__).resolve().parent.parent / "UPLOAD_READY_R1_FINAL6" / "SuppTables_R1.xlsx"
OUT_PDF = ROOT / "Supplementary_Figure_S3_animal_enrichment_summary.pdf"
OUT_PNG = ROOT / "Supplementary_Figure_S3_animal_enrichment_summary.png"

BLUE = "#2468A2"
RED = "#B63A3A"
GREY = "#58646F"


def rows(sheet):
    wb = load_workbook(BOOK, read_only=True, data_only=True)
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def wrap(value, width=34):
    return "\n".join(textwrap.wrap(str(value), width=width, break_long_words=False))


def panel_label(ax, label):
    ax.text(-0.08, 1.05, label, transform=ax.transAxes, fontsize=13,
            fontweight="bold", va="top")


def main():
    strata = [
        ("S2_GO_uremic", "Uraemic calcification · GO BP", BLUE),
        ("S2_KEGG_uremic", "Uraemic calcification · KEGG", BLUE),
        ("S2_GO_AAD", "AAD lesions · GO BP", RED),
        ("S2_KEGG_AAD", "AAD lesions · KEGG", RED),
    ]

    ora = []
    for sheet, label, colour in strata:
        sig = [r for r in rows(sheet) if r["p.adjust"] is not None and r["p.adjust"] < 0.05]
        chosen = sorted(sig, key=lambda r: (r["p.adjust"], str(r["ID"])))[:5]
        for rank, r in enumerate(chosen, 1):
            ora.append({**r, "stratum": label, "colour": colour, "rank": rank})

    go = []
    for sheet, label, colour in [strata[0], strata[2]]:
        sig = [r for r in rows(sheet) if r["p.adjust"] is not None and r["p.adjust"] < 0.05]
        for rank, r in enumerate(sorted(sig, key=lambda r: (r["p.adjust"], str(r["ID"])))[:10], 1):
            go.append({**r, "stratum": label.replace(" · GO BP", ""), "colour": colour, "rank": rank})

    gsea_all = rows("S2_GSEA_all")
    selected = []
    for dataset, db, title, colour in [
        ("Uremic", "GO_BP", "Uraemic calcification · GO BP", BLUE),
        ("Uremic", "KEGG", "Uraemic calcification · KEGG", BLUE),
        ("AAD", "GO_BP", "AAD lesions · GO BP", RED),
        ("AAD", "KEGG", "AAD lesions · KEGG", RED),
    ]:
        sig = [r for r in gsea_all if r["dataset"] == dataset and r["database"] == db
               and r["p.adjust"] is not None and r["p.adjust"] < 0.05]
        pick = sorted(sig, key=lambda r: (r["p.adjust"], -abs(r["NES"]), str(r["ID"])))[0]
        selected.append(({**pick, "title": title, "colour": colour}))

    # Two publication-width pages preserve all six panels without shrinking text.
    plt.rcParams.update({'font.size':7, 'axes.labelsize':7.5, 'xtick.labelsize':7,
                         'ytick.labelsize':7, 'pdf.fonttype':42})
    fig = plt.figure(figsize=(183/25.4, 275/25.4), constrained_layout=True)
    gs = fig.add_gridspec(2, 1, height_ratios=[1.05, 1.2])
    ax_a = fig.add_subplot(gs[0, 0])
    ax_b = fig.add_subplot(gs[1, 0])
    fig_detail, axs = plt.subplots(2, 2, figsize=(183/25.4, 150/25.4), constrained_layout=True)
    mini = list(axs.flat)

    # A: five most significant ORA terms per model/database.
    display = list(reversed(ora))
    y = list(range(len(display)))
    logq = [-math.log10(max(r["p.adjust"], 1e-300)) for r in display]
    norm = Normalize(vmin=min(logq), vmax=max(logq))
    cmap = plt.get_cmap("viridis")
    sizes = [0.35 * r["Count"] for r in display]
    ax_a.scatter([r["FoldEnrichment"] for r in display], y, s=sizes,
                 c=logq, cmap=cmap, norm=norm, edgecolor="white", linewidth=0.4)
    ax_a.set_yticks(y, [wrap(f'{r["stratum"]}: {r["Description"]}', 55) for r in display], fontsize=6.5)
    ax_a.set_xlabel("Fold enrichment")
    ax_a.set_title("Top five significant ORA terms per model and database", loc="left", fontsize=8)
    ax_a.grid(axis="x", color="#DDE2E6", linewidth=0.6)
    sm = ScalarMappable(norm=norm, cmap=cmap); sm.set_array([])
    cb = fig.colorbar(sm, ax=ax_a, fraction=0.03, pad=0.02)
    cb.set_label("−log10(BH-adjusted P)", fontsize=8)
    handles = [ax_a.scatter([],[],s=0.35*n,color='#777777',label=str(n)) for n in (20,50,100)]
    ax_a.legend(handles=handles,title='Mapped genes',fontsize=6.5,title_fontsize=6.5,
                loc='upper right',frameon=False)
    panel_label(ax_a, "A")

    # B: ten most significant GO BP ORA terms for each model.
    renal = [r for r in go if r["stratum"] == "Uraemic calcification"]
    aad = [r for r in go if r["stratum"] == "AAD lesions"]
    assert len(renal) == len(aad) == 10, "Each model must retain ten frozen GO terms"
    renal = list(reversed(renal)); aad = list(reversed(aad))
    labels = [wrap(r["Description"], 34) for r in renal] + [wrap(r["Description"], 34) for r in aad]
    vals = [r["FoldEnrichment"] for r in renal + aad]
    cols = [BLUE] * len(renal) + [RED] * len(aad)
    yy = list(range(len(labels)))
    ax_b.barh(yy, vals, color=cols, height=0.72)
    ax_b.set_yticks(yy, labels, fontsize=6.5)
    ax_b.set_xlabel("Fold enrichment\nBlue: uraemic calcification; red: AAD lesions", fontsize=8)
    ax_b.set_title("Top ten GO BP ORA terms per model", loc="left", fontsize=8)
    ax_b.grid(axis="x", color="#DDE2E6", linewidth=0.6)
    panel_label(ax_b, "B")

    # C-F: deterministically selected all-ranked GSEA summaries.
    axes = mini
    for ax, label, r in zip(axes, list("CDEF"), selected):
        nes = float(r["NES"])
        ax.axvline(0, color="#777777", linewidth=0.8)
        ax.barh([0], [nes], color=r["colour"], height=0.24)
        ax.set_xlim(min(-0.4, nes - 0.45), max(0.4, nes + 0.45))
        ax.set_ylim(-1.2, 1.2)
        ax.set_yticks([])
        ax.set_xlabel("NES", fontsize=8)
        ax.set_title(r["title"].replace(' · ','\n'), loc="left", fontsize=8)
        ax.text(0.02, 0.90, wrap(r["Description"], 32), transform=ax.transAxes,
                ha="left", va="top", fontsize=8.0, fontweight="bold")
        ax.text(0.02, 0.34,
                f'ID: {r["ID"]}\nNES = {nes:.2f}\nBH-adjusted P = {r["p.adjust"]:.2g}\nGene-set size = {int(r["setSize"])}',
                transform=ax.transAxes, ha="left", va="top", fontsize=7.5, color=GREY)
        ax.grid(axis="x", color="#E3E7EA", linewidth=0.6)
        panel_label(ax, label)

    fig.suptitle("Supplementary Figure S3 | Animal enrichment summaries (1/2)",
                 fontsize=9, fontweight="bold")
    fig_detail.suptitle("Supplementary Figure S3 | All-ranked summaries (2/2)",
                        fontsize=9,fontweight='bold')
    with PdfPages(OUT_PDF) as pdf:
        pdf.savefig(fig);pdf.savefig(fig_detail)
    fig.savefig(OUT_PNG,dpi=600)
    fig_detail.savefig(OUT_PNG.with_name(OUT_PNG.stem+'_page2.png'),dpi=600)
    payload={'source_workbook_sha256':hashlib.sha256(BOOK.read_bytes()).hexdigest(),
             'ora':ora,'go':go,'gsea_selected':selected,
             'page_width_mm':183,'page_heights_mm':[275,150],
             'point_area_mapping':'0.35 * mapped-gene count (points squared)',
             'new_analysis':False}
    (ROOT/'S3_publication_layout_source.json').write_text(json.dumps(payload,indent=2,default=str))
    print(OUT_PDF)
    print(OUT_PNG)


if __name__ == "__main__":
    main()
