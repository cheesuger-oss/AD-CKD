#!/usr/bin/env python3
"""Build auditable Phase 1 expression matrices and sample metadata."""

from __future__ import annotations

import csv
import gzip
import hashlib
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "project" / "data" / "geo"
OUT = ROOT / "project" / "results" / "geo"
PROCESSED = RAW / "processed"
DOWNLOADS = Path.home() / "Downloads"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_tsv(path: Path, rows: list[dict], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)


def audit_gse146638() -> tuple[list[dict], dict]:
    source = RAW / "GSE146638_Effect_of_CKD_on_calcified_aorta_by_RNAseq_Processed_data.xlsx"
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb.active
    values = ws.iter_rows(values_only=True)
    header = list(next(values))
    sample_names = header[12:]
    expression_path = PROCESSED / "GSE146638_expression.tsv.gz"
    genes: list[str] = []
    missing = 0
    with gzip.open(expression_path, "wt", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["gene_symbol", *sample_names])
        for row in values:
            genes.append(str(row[0]))
            vals = list(row[12:])
            missing += sum(v is None for v in vals)
            writer.writerow([row[0], *vals])

    metadata = []
    gsm = [f"GSM{n}" for n in range(4403078, 4403092)]
    for i, (sample, accession) in enumerate(zip(sample_names, gsm, strict=True)):
        if i < 5:
            group, include = "Control", "yes"
        elif i < 10:
            group, include = "Uremic", "yes"
        else:
            group, include = "Control_VitD", "no"
        metadata.append(
            {
                "dataset": "GSE146638",
                "expression_column": sample,
                "geo_accession": accession,
                "species": "Rattus norvegicus",
                "sex": "not_reported",
                "treatment": group,
                "pathology": group,
                "primary_group": group,
                "include_primary": include,
                "exclusion_reason": "" if include == "yes" else "vitamin D group excluded from primary contrast",
            }
        )
    dup = sum(v - 1 for v in Counter(genes).values() if v > 1)
    stats = {
        "source": str(source.relative_to(ROOT)),
        "sha256": sha256(source),
        "features": len(genes),
        "samples": len(sample_names),
        "missing_expression_cells": missing,
        "duplicate_gene_rows": dup,
        "identifier": "rat gene symbol (submitter-provided)",
        "scale": "processed non-integer expression; analyzed after log2(x+1)",
    }
    return metadata, stats


def parse_soft(path: Path) -> dict[str, dict[str, str]]:
    samples: dict[str, dict[str, str]] = {}
    current: dict[str, str] | None = None
    with gzip.open(path, "rt", encoding="utf-8", errors="replace") as handle:
        for line in handle:
            line = line.rstrip("\r\n")
            if line.startswith("^SAMPLE = "):
                accession = line.split(" = ", 1)[1]
                current = {"geo_accession": accession}
            elif current is not None and line.startswith("!Sample_title = "):
                current["title"] = line.split(" = ", 1)[1]
                samples[current["title"]] = current
            elif current is not None and line.startswith("!Sample_characteristics_ch1 = Sex: "):
                current["sex"] = line.rsplit(": ", 1)[1]
            elif current is not None and line.startswith("!Sample_characteristics_ch1 = treatment: "):
                current["treatment"] = line.split("treatment: ", 1)[1]
            elif current is not None and line.startswith("!Sample_description = Column name in raw_counts.txt:"):
                current["expression_column"] = line.split(":", 1)[1]
    return samples


def classify_mouse(name: str) -> tuple[str, str, str]:
    sex = "Male" if name.startswith("Ma") else "Female"
    if "NonAD" in name:
        return sex, "BAPN", "NonAD"
    if "AAD" in name:
        return sex, "BAPN", "AAD"
    return sex, "Control", "Control"


def audit_gse325608() -> tuple[list[dict], dict]:
    source = RAW / "GSE325608_raw_counts.txt.gz"
    soft_candidates = [DOWNLOADS / "GSE325608_family.soft.gz", Path("/tmp/GSE325608_family.soft.gz")]
    soft = next((p for p in soft_candidates if p.exists()), None)
    if soft is None:
        raise FileNotFoundError("GSE325608_family.soft.gz not found in Downloads")
    official = parse_soft(soft)
    with gzip.open(source, "rt", encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        header = next(reader)
        raw_columns = header[1:]
        names = [re.sub(r"^read_count_", "", x) for x in raw_columns]
        genes: list[str] = []
        missing = noninteger = negative = zero = total = 0
        for row in reader:
            genes.append(row[0])
            for x in row[1:]:
                total += 1
                if x == "":
                    missing += 1
                    continue
                v = float(x)
                noninteger += v != int(v)
                negative += v < 0
                zero += v == 0

    if set(names) != set(official):
        raise ValueError(f"SOFT/expression sample mismatch: expression-only={set(names)-set(official)}, soft-only={set(official)-set(names)}")
    metadata = []
    for raw, name in zip(raw_columns, names, strict=True):
        sex, treatment, pathology = classify_mouse(name)
        record = official[name]
        if record.get("sex") != sex:
            raise ValueError(f"Sex mismatch for {name}")
        include = pathology in {"Control", "AAD"}
        metadata.append(
            {
                "dataset": "GSE325608",
                "expression_column": raw,
                "geo_accession": record["geo_accession"],
                "species": "Mus musculus",
                "sex": sex,
                "treatment": treatment,
                "pathology": pathology,
                "primary_group": pathology,
                "include_primary": "yes" if include else "no",
                "exclusion_reason": "" if include else "BAPN-exposed without AD; retained for sensitivity analysis only",
            }
        )
    dup = sum(v - 1 for v in Counter(genes).values() if v > 1)
    stats = {
        "source": str(source.relative_to(ROOT)),
        "sha256": sha256(source),
        "soft_source": str(soft),
        "soft_sha256": sha256(soft),
        "features": len(genes),
        "samples": len(names),
        "missing_expression_cells": missing,
        "duplicate_gene_rows": dup,
        "noninteger_cells": noninteger,
        "negative_cells": negative,
        "zero_fraction": zero / total,
        "identifier": "mouse gene symbol (submitter-provided; GRCm39)",
        "scale": "RSEM estimated counts (nonnegative, partly fractional); limma-voom, not DESeq2",
    }
    return metadata, stats


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    PROCESSED.mkdir(parents=True, exist_ok=True)
    meta1, stats1 = audit_gse146638()
    meta2, stats2 = audit_gse325608()
    fields = list(meta1[0])
    write_tsv(OUT / "GSE146638_sample_metadata.tsv", meta1, fields)
    write_tsv(OUT / "GSE325608_sample_metadata.tsv", meta2, fields)
    write_tsv(OUT / "phase1_sample_metadata.tsv", meta1 + meta2, fields)

    counts1 = Counter(x["primary_group"] for x in meta1 if x["include_primary"] == "yes")
    counts2 = Counter(x["primary_group"] for x in meta2 if x["include_primary"] == "yes")
    all2 = Counter(x["primary_group"] for x in meta2)
    report = f"""# Phase 1 数据与元数据审计

生成日期：2026-08-07

## 冻结主对比

- GSE146638：Uremic vs Control，{counts1['Uremic']} vs {counts1['Control']}；4 个 Control+VitD 样本排除。
- GSE325608：AAD vs Control，{counts2['AAD']} vs {counts2['Control']}；主模型校正 sex；{all2['NonAD']} 个 BAPN-NonAD 样本仅用于敏感性分析。

## 数据结构

- GSE146638：{stats1['features']} 个基因、{stats1['samples']} 个样本、缺失 {stats1['missing_expression_cells']}、重复 symbol 行 {stats1['duplicate_gene_rows']}。输入是已处理的连续表达值，主分析采用 limma（log2(x+1)）。
- GSE325608：{stats2['features']} 个基因、{stats2['samples']} 个样本、缺失 {stats2['missing_expression_cells']}、重复 symbol 行 {stats2['duplicate_gene_rows']}、负值 {stats2['negative_cells']}、零值比例 {stats2['zero_fraction']:.3f}。RSEM v1.3.1 生成的 estimated counts 含 {stats2['noninteger_cells']} 个非整数单元格，采用 limma-voom，避免不当使用 DESeq2 整数计数似然。

## 标识和映射规则

- GSE146638：提交者提供的大鼠 gene symbol；重复 symbol 在 log2 转换前按原始表达求和。
- GSE325608：GRCm39 小鼠 gene symbol；当前无重复 symbol。
- 跨物种：Ensembl BioMart，仅保留 `ortholog_one2one` 且 orthology confidence=1 的大鼠/小鼠→人类映射；保存完整映射表和未映射记录。

## 已解决的待确认项

- GSE325608 的正式样本组成是 Control=20、BAPN-NonAD=12、AAD=16（男7、女9），共48例；并非把所有 BAPN 暴露动物直接合并为 AD。
- 性别与病理结局同时影响分组，主模型用 `~ sex + primary_group`；不把性别差异误作 AD 效应。
- 主证据是两个数据集的共同显著通路，不以 DEG 简单交集作为成功判据。

## 尚需在结果阶段检查

- PCA/样本相关性是否提示离群样本或 sex×group 强交互。
- 主分析阈值下 DEG 数量是否足以支持 ORA；如不足，保留全量排序结果并用 GSEA 作敏感性证据，不事后放宽阈值冒充主分析。
"""
    (OUT / "PHASE1_METADATA_AUDIT.md").write_text(report, encoding="utf-8")
    with (OUT / "phase1_input_checksums.tsv").open("w", encoding="utf-8") as handle:
        handle.write("dataset\tfile\tsha256\n")
        for dataset, stats in (("GSE146638", stats1), ("GSE325608", stats2)):
            handle.write(f"{dataset}\t{stats['source']}\t{stats['sha256']}\n")
        handle.write(f"GSE325608\t{stats2['soft_source']}\t{stats2['soft_sha256']}\n")
    print(report)


if __name__ == "__main__":
    main()
